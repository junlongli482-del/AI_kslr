"""Small, explicitly synthetic fixtures; never presented as real publication pages."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

root=Path(__file__).resolve().parents[1]
out=root/'data/fixtures';out.mkdir(parents=True,exist_ok=True)
doc=Document()
doc.add_heading('快速录入功能验证 · 衍生测试资料',0)
doc.add_paragraph('数学基础练习（仅用于输入格式与随题答案流程测试）')
doc.add_paragraph('1. 计算：1+1=____。')
doc.add_paragraph('答案：2。')
doc.add_paragraph('解析：两个1相加，结果为2。')
doc.add_paragraph('2. 下列数中，属于偶数的是（ ）。')
doc.add_paragraph('A. 1    B. 2    C. 3    D. 5')
doc.add_paragraph('答案：B。')
doc.add_paragraph('解析：偶数能够被2整除。')
doc.save(out/'inline-answers.docx')
print('Created explicitly synthetic DOCX fixture.')
